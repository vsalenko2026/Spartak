# -*- coding: utf-8 -*-
"""Статический аудит финмодели: ищет то, на чём Excel-модели ломаются чаще
всего, не пересчитывая книгу.

    python3 tools/audit_model.py                 полный отчёт -> build/audit.md
    python3 tools/audit_model.py --kind протяжка только одну проверку
    python3 tools/audit_model.py --sheet 07_Мерч только один лист

Проверки:
  ошибка       ячейка вернула #REF!/#DIV/0!/#ЗНАЧ! и т.п.
  протяжка     в строке, протянутой по месяцам, одна-две ячейки выбиваются
               из общего шаблона — классический след ручной правки
  съезд        ссылка на вводную без $ внутри протянутой строки: при
               протяжке вправо она уезжает на соседний столбец
  хардкод      число зашито внутрь формулы вместо ссылки на 01_Вводные
  пустая       формула ссылается на пустую ячейку
  внешняя      ссылка на другой файл
  цикл         ячейка входит в диапазон, который сама же суммирует
  висяк        расчётная ячейка, на которую никто не ссылается

Аудит ничего не чинит. Он выдаёт список адресов — решение по каждому
принимает человек или отдельная правка со своим замером эффекта.
"""
import argparse
import json
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import xlsx_lib as X
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
MAP = ROOT / "build" / "model_map.json"

# Числа, которым в формуле место: индексы, знаки, деления на месяцы/проценты.
BENIGN_NUMS = {0, 1, 2, 3, 4, 6, 12, 24, 30, 31, 36, 48, 60, 100, 365, 1000,
               10000, 100000, 1000000}
NUM_IN_FORMULA = re.compile(r"(?<![A-Za-zА-Яа-яЁё_$\d.])(\d+(?:\.\d+)?)(?![\d.]*[A-Za-z\d])")

# Листы-справки: там хардкод и «висяки» нормальны, это тексты и история.
INFO_SHEETS = {"00_Старт", "16_Инструкция", "16а_Логика_модели", "15_Изменения",
               "14_Ревизия", "14а_Сверка_с_ТЗ", "17_Сверка_версий",
               "11_История_25_26", "01а_Действующая_сеть", "13_Проверки"}
INPUT_SHEET = "01_Вводные"
HORIZON_COL = 61                 # BI — последний месяц горизонта (60 мес. от B)


def load_map():
    if not MAP.exists():
        sys.exit("нет build/model_map.json — сначала: python3 tools/extract_model.py")
    return json.loads(MAP.read_text(encoding="utf-8"))


def r1c1(rec):
    return X.to_r1c1(rec["f"], rec["col"], rec["row"])


def rows_of(sheet):
    by_row = defaultdict(list)
    for addr, rec in sheet["cells"].items():
        by_row[rec["row"]].append((addr, rec))
    for r in by_row:
        by_row[r].sort(key=lambda kv: kv[1]["col"])
    return by_row


class Audit:
    def __init__(self, data, only_sheet=None):
        self.data = data
        self.sheets = {k: v for k, v in data["sheets"].items()
                       if not only_sheet or k == only_sheet}
        self.findings = []
        self._hard = defaultdict(list)

    def add(self, kind, sheet, addr, text, weight=2):
        self.findings.append({"kind": kind, "sheet": sheet, "addr": addr,
                              "text": text, "weight": weight})

    # ---------------------------------------------------------- проверки ---
    def check_errors(self):
        for sname, sheet in self.sheets.items():
            for addr, rec in sheet["cells"].items():
                v = rec.get("v")
                if isinstance(v, str) and X.ERROR_RE.fullmatch(v.strip()):
                    self.add("ошибка", sname, addr,
                             f"{v} ← {rec.get('f', '')[:90]}", 1)

    def check_spread(self):
        """Разрыв протяжки: в строке доминирует один шаблон, а пара ячеек
        из него выпадает."""
        for sname, sheet in self.sheets.items():
            if sname in INFO_SHEETS:
                continue
            for r, items in rows_of(sheet).items():
                fs = [(a, rec) for a, rec in items
                      if rec["kind"] == "formula" and rec["col"] >= 2]
                if len(fs) < 8:
                    continue
                sigs = {a: r1c1(rec) for a, rec in fs}
                cnt = Counter(sigs.values())
                top, n_top = cnt.most_common(1)[0]
                if n_top < len(fs) * 0.7:
                    continue                      # строка неоднородна по замыслу
                odd = [a for a, s in sigs.items() if s != top]
                if not odd or len(odd) > max(3, len(fs) * 0.2):
                    continue
                d = dict(fs)
                cols = sorted(d[a]["col"] for a in odd)
                first_col = min(rec["col"] for _, rec in fs)
                last_col = max(rec["col"] for _, rec in fs)
                # отклонения сплошным куском в начале строки — это обычно
                # фактический период или разгон когорты, а не ошибка;
                # то же для хвоста за горизонтом модели
                # разгон в начале и технический хвост в конце — норма;
                # ошибка выглядит как отклонение в середине ряда
                head = [c for c in cols if c - first_col < len(cols)]
                head_len = 0
                while head_len < len(cols) and cols[head_len] == first_col + head_len:
                    head_len += 1
                tail_len = 0
                while tail_len < len(cols) - head_len and \
                        cols[-1 - tail_len] == last_col - tail_len:
                    tail_len += 1
                head_block = head_len + tail_len == len(cols)
                tail_block = head_block
                label = next((rec.get("v") for _, rec in items
                              if rec["kind"] == "text"), "") or ""
                addrs = sorted(odd, key=lambda x: d[x]["col"])
                self.add("протяжка", sname, addrs[0],
                         f"строка {r} «{str(label)[:40]}»: {n_top} ячеек по одному "
                         f"шаблону, выбиваются {', '.join(addrs[:6])}"
                         f"{' и др.' if len(addrs) > 6 else ''} — {d[addrs[0]]['f'][:80]}",
                         3 if (head_block or tail_block) else 1)

    def check_slip(self):
        """Ссылка на вводные без $ в протянутой строке — уедет при протяжке."""
        for sname, sheet in self.sheets.items():
            if sname in INFO_SHEETS or sname == INPUT_SHEET:
                continue
            for r, items in rows_of(sheet).items():
                fs = [(a, rec) for a, rec in items
                      if rec["kind"] == "formula" and rec["col"] >= 2]
                if len(fs) < 8:
                    continue
                for a, rec in fs:
                    body = X.strip_strings(rec["f"])
                    for m in X.REF_RE.finditer(body):
                        sh = X.unquote_sheet(m.group("sheet") or "")
                        if sh != INPUT_SHEET:
                            continue
                        if not m.group("c1").startswith("$") and \
                                not self.is_monthly_input(m.group("c1")):
                            self.add("съезд", sname, a,
                                     f"строка {r}: ссылка на вводные без $ по столбцу "
                                     f"— {m.group(0)} в {rec['f'][:70]}", 1)
                        break

    def is_monthly_input(self, ref: str):
        """Строка вводных, заполненная по месяцам (B,C,D...), — законная цель
        для ссылки без $: она и должна ехать вместе с протяжкой."""
        p = X.split_cell(ref)
        if not p:
            return False
        row = p[1]
        cells = self.data["sheets"].get(INPUT_SHEET, {}).get("cells", {})
        filled = sum(1 for a, rec in cells.items()
                     if rec["row"] == row and rec["col"] >= 2)
        return filled >= 4

    @staticmethod
    def literals(formula):
        body = X.REF_RE.sub(" ", X.strip_strings(formula))
        out = set()
        for m in NUM_IN_FORMULA.finditer(body):
            try:
                val = float(m.group(1))
            except ValueError:
                continue
            # Хардкод, который стоит искать, — это экономический параметр:
            # цена, сумма, ставка. Мелкие целые (5, 12, 19) в этой модели —
            # индексы когорт и возрастов, их выносить некуда.
            if val in BENIGN_NUMS:
                continue
            if val.is_integer() and val < 1000:
                continue
            out.add(m.group(1))
        return out

    def check_hardcode(self):
        """Хардкод — это константа, одинаковая по всей протянутой строке.
        Число, которое от ячейки к ячейке растёт (смещение месяца в DATE,
        возраст когорты), — часть протяжки, а не зашитый параметр."""
        for sname, sheet in self.sheets.items():
            if sname in INFO_SHEETS or sname == INPUT_SHEET:
                continue
            for r, items in rows_of(sheet).items():
                fs = [(a, rec) for a, rec in items if rec["kind"] == "formula"]
                if not fs:
                    continue
                per_cell = {a: self.literals(rec["f"]) for a, rec in fs}
                if len(fs) >= 8:
                    common = set.intersection(*per_cell.values()) if per_cell else set()
                else:
                    common = set().union(*per_cell.values()) if per_cell else set()
                for b in common:
                    a0, rec0 = fs[0]
                    self._hard[b].append((sname, a0, rec0["f"]))

    def flush_hardcode(self):
        """Одно замечание на число, а не на каждое вхождение: 20 тысяч адресов
        не читает никто, а полсотни чисел — вполне."""
        for value, hits in sorted(self._hard.items(), key=lambda kv: -len(kv[1])):
            sheets = sorted({s for s, _, _ in hits})
            where = ", ".join(f"{s}!{a}" for s, a, _ in hits[:3])
            self.add("хардкод", sheets[0], hits[0][1],
                     f"число {value} зашито в формулы {len(hits)} раз "
                     f"({', '.join(sheets[:4])}): {where}"
                     f"{', ...' if len(hits) > 3 else ''} — "
                     f"{hits[0][2][:70]}",
                     1 if len(hits) >= 50 else 2)

    def check_horizon(self):
        """Горизонт модели — 60 месяцев, B:BI (столбцы 2..61). Часть листов
        считает дальше, до BO: это технический задел под «целевой запас на
        следующие 6 месяцев». Не ошибка, но знать про него надо — иначе
        следующая правка примет хвост за мусор и снесёт его."""
        for sname, sheet in self.sheets.items():
            if sname in INFO_SHEETS:
                continue
            beyond = defaultdict(list)
            for addr, rec in sheet["cells"].items():
                if rec["kind"] == "formula" and rec["col"] > HORIZON_COL:
                    beyond[rec["row"]].append(addr)
            if not beyond:
                continue
            n = sum(len(v) for v in beyond.values())
            cols = sorted({X.split_cell(a)[0] for addrs in beyond.values() for a in addrs},
                          key=lambda c: len(c))
            self.add("горизонт", sname, sorted(next(iter(beyond.values())))[0],
                     f"{n} формул в {len(beyond)} строках считаются за горизонтом "
                     f"60 мес. (столбцы {cols[0]}…{cols[-1]})", 3)

    def check_empty_refs(self):
        for sname, sheet in self.sheets.items():
            if sname in INFO_SHEETS:
                continue
            for addr, rec in sheet["cells"].items():
                if rec["kind"] != "formula":
                    continue
                for ref in rec.get("refs", []):
                    if ref["kind"] != "cell":
                        continue
                    tgt = self.data["sheets"].get(ref["sheet"])
                    if tgt and ref["ref"] not in tgt["cells"]:
                        self.add("пустая", sname, addr,
                                 f"ссылается на пустую {ref['sheet']}!{ref['ref']} "
                                 f"в {rec['f'][:70]}", 2)
                        break

    def check_external(self):
        for sname, sheet in self.sheets.items():
            for addr, rec in sheet["cells"].items():
                if rec["kind"] == "formula" and "[" in rec["f"]:
                    self.add("внешняя", sname, addr, rec["f"][:110], 1)

    def check_selfsum(self):
        for sname, sheet in self.sheets.items():
            for addr, rec in sheet["cells"].items():
                if rec["kind"] != "formula":
                    continue
                col = get_column_letter(rec["col"])
                for ref in rec.get("refs", []):
                    if ref["sheet"] != sname or ref["kind"] == "cell":
                        continue
                    if X.covers(ref["ref"], col, rec["row"],
                                sheet["max_row"], sheet["max_col"]):
                        self.add("цикл", sname, addr,
                                 f"входит в собственный диапазон {ref['ref']}: "
                                 f"{rec['f'][:80]}", 1)
                        break

    def covered_cells(self):
        """Для каждого листа — множество ячеек, на которые хоть кто-то
        ссылается. Ссылки дедуплицируются и разворачиваются один раз:
        иначе проверка каждой формулы против каждой ссылки не считается."""
        refs = defaultdict(set)
        for sheet in self.data["sheets"].values():
            for rec in sheet["cells"].values():
                if rec["kind"] != "formula":
                    continue
                for ref in rec.get("refs", []):
                    refs[ref["sheet"]].add(ref["ref"])
        covered = {}
        for sname, rs in refs.items():
            sheet = self.data["sheets"].get(sname)
            if not sheet:
                continue
            hit = set()
            for rf in rs:
                hit.update(X.expand(rf, sheet["max_row"], sheet["max_col"]))
            covered[sname] = hit
        return covered

    def check_orphans(self):
        covered = self.covered_cells()
        for sname, sheet in self.sheets.items():
            if sname in INFO_SHEETS or sname == "02_Модель":
                continue
            hit = covered.get(sname, set())
            orphan_rows = defaultdict(list)
            for addr, rec in sheet["cells"].items():
                if rec["kind"] != "formula":
                    continue
                v = rec.get("v")
                if not isinstance(v, (int, float)) or v == 0:
                    continue
                if (get_column_letter(rec["col"]), rec["row"]) not in hit:
                    orphan_rows[rec["row"]].append(addr)
            for r, addrs in sorted(orphan_rows.items()):
                label = ""
                for a, rec in sheet["cells"].items():
                    if rec["row"] == r and rec["kind"] == "text" and rec["col"] <= 2:
                        label = str(rec.get("v", ""))[:40]
                        break
                self.add("висяк", sname, sorted(addrs)[0],
                         f"строка {r} «{label}»: {len(addrs)} ячеек считаются, "
                         f"но никем не используются ({', '.join(sorted(addrs)[:4])}"
                         f"{', ...' if len(addrs) > 4 else ''})", 3)

    def run(self, kinds=None):
        all_checks = {
            "ошибка": self.check_errors,
            "протяжка": self.check_spread,
            "съезд": self.check_slip,
            "хардкод": lambda: (self.check_hardcode(), self.flush_hardcode()),
            "горизонт": self.check_horizon,
            "пустая": self.check_empty_refs,
            "внешняя": self.check_external,
            "цикл": self.check_selfsum,
            "висяк": self.check_orphans,
        }
        for name, fn in all_checks.items():
            if kinds and name not in kinds:
                continue
            fn()
        return self.findings


PRIORITY = {1: "важно", 2: "стоит посмотреть", 3: "к сведению"}


def report(findings, path: Path):
    by_kind = defaultdict(list)
    for f in findings:
        by_kind[f["kind"]].append(f)
    lines = ["# Аудит финмодели", "",
             f"Найдено {len(findings)} замечаний.", "",
             "| Проверка | Замечаний | Важно |", "|---|---:|---:|"]
    for kind in sorted(by_kind, key=lambda k: -len(by_kind[k])):
        items = by_kind[kind]
        lines.append(f"| {kind} | {len(items)} | "
                     f"{sum(1 for i in items if i['weight'] == 1)} |")
    for kind in sorted(by_kind, key=lambda k: min(i["weight"] for i in by_kind[k])):
        items = sorted(by_kind[kind], key=lambda i: (i["weight"], i["sheet"], i["addr"]))
        lines += ["", f"## {kind} — {len(items)}", ""]
        for i in items[:400]:
            lines.append(f"- **{i['sheet']}!{i['addr']}** ({PRIORITY[i['weight']]}) — {i['text']}")
        if len(items) > 400:
            lines.append(f"- ... ещё {len(items) - 400}")
    path.write_text("\n".join(lines), encoding="utf-8")


def main():
    ap = argparse.ArgumentParser(description="Статический аудит финмодели")
    ap.add_argument("--kind", action="append", help="только эти проверки")
    ap.add_argument("--sheet", help="только этот лист")
    ap.add_argument("--show", type=int, default=12, help="сколько показать в консоли")
    a = ap.parse_args()

    data = load_map()
    audit = Audit(data, a.sheet)
    findings = audit.run(set(a.kind) if a.kind else None)

    by_kind = defaultdict(list)
    for f in findings:
        by_kind[f["kind"]].append(f)
    print(f"{'проверка':<12}{'всего':>8}{'важно':>8}")
    for kind in sorted(by_kind, key=lambda k: -len(by_kind[k])):
        items = by_kind[kind]
        print(f"{kind:<12}{len(items):>8}{sum(1 for i in items if i['weight'] == 1):>8}")
    print(f"{'ИТОГО':<12}{len(findings):>8}"
          f"{sum(1 for i in findings if i['weight'] == 1):>8}")

    top = sorted(findings, key=lambda i: i["weight"])[:a.show]
    if top:
        print("\nсамое важное:")
        for i in top:
            print(f"  [{i['kind']}] {i['sheet']}!{i['addr']}  {i['text'][:110]}")

    out = ROOT / "build" / "audit.md"
    report(findings, out)
    print(f"\nотчёт: {out.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
