# -*- coding: utf-8 -*-
"""Строительные блоки книги V9.0: стили, раскладка, запись строк.

Раскладка каждого расчётного листа одна и та же:
    A  — номер блока / номер строки расчёта
    B  — название показателя
    C  — единица измерения
    D  — откуда берётся (короткое пояснение формулы словами)
    E:I — итоги по годам 1..5
    J   — итого за горизонт
    L:  — 60 месяцев

Годы слева, месяцы справа: руководитель видит пятилетку, не прокручивая,
а аналитик уходит вправо за помесячной раскладкой.
"""
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L

RED = "9E1B31"
GREY = "F2F2F2"
YELLOW = "FFF2CC"
BLUE = "DDEBF7"
GREEN = "E2F0D9"
WHITE = "FFFFFF"

КОЛ_ГОД1 = 5           # E
КОЛ_ИТОГО = 10         # J
КОЛ_МЕС1 = 12          # L
ЛЕТ = 5
МЕСЯЦЕВ = 60            # горизонт, который показываем и суммируем
ХВОСТ = 6               # технические месяцы за горизонтом: по ним считается
                        # целевой запас склада на будущие продажи
ВСЕГО_МЕСЯЦЕВ = МЕСЯЦЕВ + ХВОСТ

# Коды форматов Excel хранит всегда в канонической записи: запятая —
# разделитель разрядов, точка — десятичный. Русские «# ##0,0» Excel читает
# как «# ##0» с делением на тысячу и рисует 6 200 школ как «06».
# Разделитель разрядов ставим деньгам, штукам и людям; процентам,
# коэффициентам и индексам — нет, они меньше тысячи по смыслу.
ДЕНЬГИ = '#,##0;[Red]-#,##0;-'
ДЕНЬГИ_МЛН = '#,##0.0;[Red]-#,##0.0;-'
ШТУКИ = '#,##0.0;[Red]-#,##0.0;-'
ЦЕЛЫЕ = '#,##0;[Red]-#,##0;-'
ПРОЦЕНТ = '0.0%;[Red]-0.0%;-'
КОЭФФ = '0.000;[Red]-0.000;-'
ДАТА = 'mmm yyyy'


def мес(i):
    """Буква столбца месяца i (1..60) на новых листах."""
    return L(КОЛ_МЕС1 + i - 1)


def тариф(i):
    """Буква столбца месяца i на перенесённых листах (07_Тарифы): там
    месяцы начинаются со столбца B, а не с L."""
    return L(1 + i)


def год(y):
    """Буква столбца года y (1..5)."""
    return L(КОЛ_ГОД1 + y - 1)


class Лист:
    """Обёртка над листом: пишет строки расчёта в единой раскладке."""

    def __init__(self, ws, заголовок, подзаголовок=""):
        self.ws = ws
        self.строка = 1
        ws.sheet_view.showGridLines = False
        ws["B1"] = заголовок
        ws["B1"].font = Font(bold=True, size=14, color=RED)
        if подзаголовок:
            ws["B2"] = подзаголовок
            ws["B2"].font = Font(italic=True, size=10, color="595959")
        self.строка = 4
        self._ширины()

    def _ширины(self):
        ws = self.ws
        ws.column_dimensions["A"].width = 5
        ws.column_dimensions["B"].width = 44
        ws.column_dimensions["C"].width = 11
        ws.column_dimensions["D"].width = 46
        for y in range(ЛЕТ):
            ws.column_dimensions[L(КОЛ_ГОД1 + y)].width = 15
        ws.column_dimensions[L(КОЛ_ИТОГО)].width = 16
        ws.column_dimensions[L(КОЛ_ИТОГО + 1)].width = 3
        for i in range(1, ВСЕГО_МЕСЯЦЕВ + 1):
            ws.column_dimensions[мес(i)].width = 13
            if i > МЕСЯЦЕВ:                      # хвост скрыт: он технический
                ws.column_dimensions[мес(i)].hidden = True

    def шапка(self, дата_строка=None):
        """Две строки заголовков: номер месяца и под ним календарная дата.

        «М14» само по себе ничего не говорит; календарная подпись снимает
        вопрос, какой это месяц и год. Дата — настоящая дата с форматом, а
        не текст: так подпись не зависит от языка Excel."""
        ws, r = self.ws, self.строка
        rд = r + 1
        ws[f"B{r}"] = "Показатель"
        ws[f"C{r}"] = "Ед."
        ws[f"D{r}"] = "Как считается"
        for y in range(1, ЛЕТ + 1):
            ws[f"{год(y)}{r}"] = f"Год {y}"
        ws[f"{L(КОЛ_ИТОГО)}{r}"] = "Итого 5 лет"
        for i in range(1, ВСЕГО_МЕСЯЦЕВ + 1):
            c = мес(i)
            ws[f"{c}{r}"] = (f"М{i}" if i <= МЕСЯЦЕВ
                             else f"тех. {i - МЕСЯЦЕВ}")
            ws[f"{c}{rд}"] = ("=DATE(YEAR(СтартМодели),"
                              f"MONTH(СтартМодели)+{i}-1,1)")
            ws[f"{c}{rд}"].number_format = ДАТА
        # подписи слева от месяцев занимают обе строки шапки
        for col in range(2, КОЛ_ИТОГО + 1):
            ws.merge_cells(start_row=r, start_column=col,
                           end_row=rд, end_column=col)
        for строка_шапки in (r, rд):
            for col in range(2, КОЛ_МЕС1 + ВСЕГО_МЕСЯЦЕВ):
                c = ws.cell(row=строка_шапки, column=col)
                c.font = Font(bold=True, color=WHITE, size=10)
                c.fill = PatternFill("solid", fgColor=RED)
                c.alignment = Alignment(horizontal="center", vertical="center",
                                        wrap_text=True)
        ws.freeze_panes = f"{L(КОЛ_МЕС1)}{rд + 1}"
        ws.row_dimensions[r].height = 26
        ws.row_dimensions[rд].height = 16
        self.строка += 2
        return r

    def блок(self, название):
        ws, r = self.ws, self.строка
        ws[f"B{r}"] = название
        ws[f"B{r}"].font = Font(bold=True, size=11, color=RED)
        ws[f"B{r}"].fill = PatternFill("solid", fgColor=YELLOW)
        for col in range(2, КОЛ_МЕС1 + ВСЕГО_МЕСЯЦЕВ):
            ws.cell(row=r, column=col).fill = PatternFill("solid", fgColor=YELLOW)
        self.строка += 1
        return r

    # Что именно стоит в колонках лет, если это не сумма за год. Подпись
    # шапки «Итого 5 лет» верна только для потоков; для остатков там снимок
    # последнего месяца, и это надо говорить вслух, а не оставлять читателю.
    ПОЯСНЕНИЕ_ИТОГА = {
        "конец": "В колонках лет — значение на конец года, в «Итого» — на "
                 "конец пятого года. Это остаток, а не сумма.",
        "минимум": "В колонках лет — минимум за год, в «Итого» — минимум за "
                   "все 60 месяцев.",
        "среднее": "В колонках лет — среднее за год, в «Итого» — среднее за "
                   "все 60 месяцев.",
    }

    def ряд(self, название, формула, ед="₽", как="", формат=ДЕНЬГИ,
            жирный=False, первый=None, итог="сумма", годы=True):
        """Строка расчёта: месяцы считаются формулой, годы — сводкой.

        `формула` — шаблон с {м} вместо буквы столбца месяца и {пред} вместо
        предыдущего. `первый` — отдельная формула для первого месяца.
        `годы=False` оставляет колонки лет пустыми: так помечаются строки
        внутренней механики, у которых годового смысла нет."""
        ws, r = self.ws, self.строка
        # Людей дробными не считают: «4 119 детей» читается, «4 118,9» — нет.
        # Школы остаются с десятой: когорта из половины школы — это правда
        # модели, округление до целой сломало бы сходимость сумм.
        if ед == "чел." and формат == ШТУКИ:
            формат = ЦЕЛЫЕ
        ws[f"B{r}"] = название
        ws[f"C{r}"] = ед
        ws[f"D{r}"] = как
        ws[f"D{r}"].font = Font(italic=True, size=9, color="595959")
        ws[f"D{r}"].alignment = Alignment(wrap_text=True, vertical="top")
        for i in range(1, ВСЕГО_МЕСЯЦЕВ + 1):
            c = мес(i)
            ш = первый if (i == 1 and первый is not None) else формула
            if ш is None:
                continue
            знач = ш.format(м=c, пред=мес(i - 1) if i > 1 else None,
                            пред2=мес(i - 2) if i > 2 else None,
                            пред3=мес(i - 3) if i > 3 else None,
                            **{f"след{k}": (мес(i + k) if i + k <= ВСЕГО_МЕСЯЦЕВ
                                            else мес(ВСЕГО_МЕСЯЦЕВ))
                               for k in range(1, 7)},
                            т=тариф(i), тпред=тариф(i - 1) if i > 1 else None,
                            i=i, r=r)
            ws[f"{c}{r}"] = знач if str(знач).startswith("=") else знач
            ws[f"{c}{r}"].number_format = формат
        if как:
            хвост = self.ПОЯСНЕНИЕ_ИТОГА.get(итог, "")
            if хвост:
                ws[f"D{r}"] = f"{как} {хвост}"
        elif итог in self.ПОЯСНЕНИЕ_ИТОГА:
            ws[f"D{r}"] = self.ПОЯСНЕНИЕ_ИТОГА[итог]
        if not годы:
            ws[f"D{r}"] = (как or "") + (" " if как else "") + \
                "Внутренняя механика расчёта: годового итога у этой строки " \
                "нет, смотрите итоговые строки блока ниже."
            ws[f"D{r}"].font = Font(italic=True, size=9, color="595959")
            ws[f"D{r}"].alignment = Alignment(wrap_text=True, vertical="top")
            for y in range(ЛЕТ):
                ws.cell(row=r, column=КОЛ_ГОД1 + y).fill = \
                    PatternFill("solid", fgColor=GREY)
            if жирный:
                for col in list(range(2, 5)) + \
                           list(range(КОЛ_МЕС1, КОЛ_МЕС1 + ВСЕГО_МЕСЯЦЕВ)):
                    ws.cell(row=r, column=col).font = Font(bold=True)
            self.строка += 1
            return r
        # годы и итог
        for y in range(1, ЛЕТ + 1):
            a, b = мес((y - 1) * 12 + 1), мес(y * 12)
            ф = (f"=SUM({a}{r}:{b}{r})" if итог == "сумма"
                 else f"={b}{r}" if итог == "конец"
                 else f"=MIN({a}{r}:{b}{r})" if итог == "минимум"
                 else f"=AVERAGE({a}{r}:{b}{r})")
            ws[f"{год(y)}{r}"] = ф
            ws[f"{год(y)}{r}"].number_format = формат
        итого = (f"=SUM({мес(1)}{r}:{мес(МЕСЯЦЕВ)}{r})" if итог == "сумма"
                 else f"={мес(МЕСЯЦЕВ)}{r}" if итог == "конец"
                 else f"=MIN({мес(1)}{r}:{мес(МЕСЯЦЕВ)}{r})" if итог == "минимум"
                 else f"=AVERAGE({мес(1)}{r}:{мес(МЕСЯЦЕВ)}{r})")
        ws[f"{L(КОЛ_ИТОГО)}{r}"] = итого
        ws[f"{L(КОЛ_ИТОГО)}{r}"].number_format = формат
        if жирный:
            for col in list(range(2, КОЛ_ИТОГО + 1)) + \
                       list(range(КОЛ_МЕС1, КОЛ_МЕС1 + ВСЕГО_МЕСЯЦЕВ)):
                ws.cell(row=r, column=col).font = Font(bold=True)
        for y in range(ЛЕТ):
            ws.cell(row=r, column=КОЛ_ГОД1 + y).fill = PatternFill("solid", fgColor=GREY)
        ws.cell(row=r, column=КОЛ_ИТОГО).fill = PatternFill("solid", fgColor=BLUE)
        self.строка += 1
        return r

    def пусто(self, n=1):
        self.строка += n
