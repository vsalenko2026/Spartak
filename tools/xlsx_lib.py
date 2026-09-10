# -*- coding: utf-8 -*-
"""Общие примитивы разбора Excel-формул для инструментов финмодели.

Здесь нет ничего про «Спартак» — только механика: разобрать формулу на
ссылки, перевести A1 в R1C1, развернуть диапазон. Всё остальное строится
поверх.
"""
import re
from openpyxl.utils import column_index_from_string, get_column_letter

# --- лексика формул -------------------------------------------------------
# Строковые литералы вырезаем до разбора: внутри них могут быть скобки,
# восклицательные знаки и что угодно ещё.
STRING_RE = re.compile(r'"(?:[^"]|"")*"')

# Ссылка: [Лист!]A1[:B2]. Имя листа либо в апострофах, либо без пробелов.
REF_RE = re.compile(
    r"(?:(?P<sheet>'(?:[^']|'')+'|[A-Za-zА-Яа-яЁё0-9_.]+)!)?"
    r"(?P<c1>\$?[A-Z]{1,3}\$?\d{1,7})"
    r"(?::(?P<c2>\$?[A-Z]{1,3}\$?\d{1,7}))?"
    r"(?![\(\w])"
)

# Ссылка на целые столбцы/строки: Лист!A:A, 5:7
BAND_RE = re.compile(
    r"(?:(?P<sheet>'(?:[^']|'')+'|[A-Za-zА-Яа-яЁё0-9_.]+)!)?"
    r"(?P<b1>\$?[A-Z]{1,3}|\$?\d{1,7}):(?P<b2>\$?[A-Z]{1,3}|\$?\d{1,7})(?![\(\w])"
)

# Имя (именованный диапазон или функция — отличаем по скобке следом).
NAME_RE = re.compile(r"(?<![\w.!$])(?P<name>[A-Za-zА-Яа-яЁё_][A-Za-zА-Яа-яЁё0-9_.]*)\s*(?P<paren>\()?")

ERROR_RE = re.compile(r"#(?:REF!|DIV/0!|VALUE!|NAME\?|N/A|NULL!|NUM!|ССЫЛКА!|ДЕЛ/0!|ЗНАЧ!|ИМЯ\?|Н/Д|ПУСТО!|ЧИСЛО!)")

CELL_RE = re.compile(r"^\$?([A-Z]{1,3})\$?(\d{1,7})$")


def strip_strings(formula: str) -> str:
    """Заменяет строковые литералы пробелами той же длины (позиции сохраняются)."""
    return STRING_RE.sub(lambda m: " " * len(m.group(0)), formula)


def unquote_sheet(name: str) -> str:
    if name and name.startswith("'") and name.endswith("'"):
        return name[1:-1].replace("''", "'")
    return name


def split_cell(ref: str):
    """'$B$7' -> ('B', 7, abs_col, abs_row)"""
    m = CELL_RE.match(ref)
    if not m:
        return None
    col, row = m.group(1), int(m.group(2))
    return col, row, ref.startswith("$"), "$" in ref[1:]


def parse_refs(formula: str, own_sheet: str):
    """Все ссылки формулы. Возвращает список dict(sheet, ref, kind)."""
    if not formula:
        return []
    body = strip_strings(formula)
    out, seen = [], set()

    for m in BAND_RE.finditer(body):
        sheet = unquote_sheet(m.group("sheet")) or own_sheet
        ref = f"{m.group('b1')}:{m.group('b2')}".replace("$", "")
        key = (sheet, ref)
        if key not in seen:
            seen.add(key)
            out.append({"sheet": sheet, "ref": ref, "kind": "band"})

    for m in REF_RE.finditer(body):
        sheet = unquote_sheet(m.group("sheet")) or own_sheet
        c1 = m.group("c1").replace("$", "")
        c2 = (m.group("c2") or "").replace("$", "")
        ref = f"{c1}:{c2}" if c2 else c1
        key = (sheet, ref)
        if key not in seen:
            seen.add(key)
            out.append({"sheet": sheet, "ref": ref, "kind": "range" if c2 else "cell"})
    return out


def parse_names(formula: str, known_names):
    """Именованные диапазоны, использованные в формуле."""
    if not formula:
        return []
    body = strip_strings(formula)
    found = []
    for m in NAME_RE.finditer(body):
        name = m.group("name")
        if m.group("paren"):          # это функция
            continue
        if name in known_names and name not in found:
            found.append(name)
    return found


def to_r1c1(formula: str, col: int, row: int) -> str:
    """A1 -> R1C1 относительно (col,row). Нужно, чтобы сравнивать соседние
    ячейки: протянутая формула в R1C1 у всех одинаковая, а разрыв виден."""
    if not formula:
        return ""
    body = formula
    out, pos = [], 0
    # строки пропускаем как есть
    protected = [(m.start(), m.end()) for m in STRING_RE.finditer(body)]

    def in_string(i):
        return any(s <= i < e for s, e in protected)

    def conv(one: str) -> str:
        p = split_cell(one)
        if not p:
            return one
        c, r, _, _ = p
        abs_col = one.startswith("$")
        rest = one[1:] if abs_col else one
        abs_row = "$" in rest
        ci = column_index_from_string(c)
        rp = f"R{r}" if abs_row else (f"R[{r - row}]" if r != row else "R")
        cp = f"C{ci}" if abs_col else (f"C[{ci - col}]" if ci != col else "C")
        return rp + cp

    for m in REF_RE.finditer(body):
        if in_string(m.start()):
            continue
        out.append(body[pos:m.start()])
        sheet = m.group("sheet")
        prefix = f"{sheet}!" if sheet else ""
        c1 = conv(m.group("c1"))
        c2 = f":{conv(m.group('c2'))}" if m.group("c2") else ""
        out.append(prefix + c1 + c2)
        pos = m.end()
    out.append(body[pos:])
    return "".join(out)


def expand(ref: str, max_row: int = 0, max_col: int = 0):
    """'B2:D3' -> [('B',2),...]. Полосы (A:A) ограничиваются размером листа."""
    if ":" not in ref:
        p = split_cell(ref)
        return [(p[0], p[1])] if p else []
    a, b = ref.split(":", 1)
    pa, pb = split_cell(a), split_cell(b)
    if not pa or not pb:                       # полоса
        if a.isdigit() and b.isdigit():
            r1, r2 = sorted((int(a), int(b)))
            return [(get_column_letter(c), r) for r in range(r1, r2 + 1)
                    for c in range(1, max(max_col, 1) + 1)]
        try:
            c1, c2 = sorted((column_index_from_string(a.replace("$", "")),
                             column_index_from_string(b.replace("$", ""))))
        except ValueError:
            return []
        return [(get_column_letter(c), r) for c in range(c1, c2 + 1)
                for r in range(1, max(max_row, 1) + 1)]
    c1, c2 = sorted((column_index_from_string(pa[0]), column_index_from_string(pb[0])))
    r1, r2 = sorted((pa[1], pb[1]))
    return [(get_column_letter(c), r) for r in range(r1, r2 + 1) for c in range(c1, c2 + 1)]


def covers(ref: str, col_letter: str, row: int, max_row: int = 0, max_col: int = 0) -> bool:
    """Попадает ли ячейка в ссылку (без разворачивания диапазона)."""
    ci = column_index_from_string(col_letter)
    if ":" not in ref:
        p = split_cell(ref)
        return bool(p) and p[0] == col_letter and p[1] == row
    a, b = ref.split(":", 1)
    pa, pb = split_cell(a), split_cell(b)
    if pa and pb:
        c1, c2 = sorted((column_index_from_string(pa[0]), column_index_from_string(pb[0])))
        r1, r2 = sorted((pa[1], pb[1]))
        return c1 <= ci <= c2 and r1 <= row <= r2
    if a.isdigit() and b.isdigit():
        r1, r2 = sorted((int(a), int(b)))
        return r1 <= row <= r2
    try:
        c1, c2 = sorted((column_index_from_string(a.replace("$", "")),
                         column_index_from_string(b.replace("$", ""))))
    except ValueError:
        return False
    return c1 <= ci <= c2
