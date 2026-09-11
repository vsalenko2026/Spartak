# -*- coding: utf-8 -*-
"""Прогон финмодели: подставить вводные, пересчитать, забрать числа.

openpyxl формулы не считает — считает LibreOffice Calc. Схема такая:
  1. книга копируется во временную папку;
  2. openpyxl подставляет вводные и стирает кэш значений;
  3. soffice --convert-to xlsx пересчитывает всю книгу;
  4. значения читаются обратно и складываются в срез (build/runs/<тег>.json).

Примеры:
    python3 tools/run_model.py                                   базовый прогон
    python3 tools/run_model.py --scenario негативный
    python3 tools/run_model.py --set '01_Вводные!B58=0.07' --tag ндс7
    python3 tools/run_model.py --compare база негативный         сравнить прогоны
"""
import argparse
import json
import re
import shutil
import subprocess
import sys
import tempfile
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
DEFAULT_MODEL = ROOT / "model" / "Spartak_V8_9_понятная.xlsx"
BUILD = ROOT / "build"
RUNS = BUILD / "runs"
RECALC = BUILD / "recalc"
LO_PROFILE = BUILD / "lo_profile"

SCENARIO_CELL = "01_Вводные!B192"
SCENARIO_NO_CELL = ("01_Вводные", "B194")     # =IF(B192=C197,1,IF(B192=D197,2,3))
# B192 хранит ТЕКСТ, а не номер: подстановка числа молча даёт позитивный
# сценарий (последняя ветка IF). Отсюда сверка номера после прогона.
SCENARIOS = {"негативный": ("Негативный", 1),
             "базовый": ("Базовый", 2),
             "позитивный": ("Позитивный", 3)}

# Вводные, которые кладём в срез: по ним тесты понимают, в каком режиме
# считалась модель, и проверяют, что режим действительно применился.
KEY_INPUTS = {
    "закрытие_в_год": ("01_Вводные", "B13"),
    "коэфф_сохранения": ("01_Вводные", "B14"),
    "лаг_открытия": ("01_Вводные", "B15"),
    "школ_на_партнёра": ("01_Вводные", "B22"),
    "целевой_запас_мес": ("01_Вводные", "B80"),
    "страховые": ("01_Вводные", "B131"),
    "каникулы_роялти_мес": ("01_Вводные", "B297"),
    "конверсия": ("01_Вводные", "B298"),
}

# что забираем из пересчитанной книги
GRAB = {
    "модель": ("02_Модель", 5, 80, 2, 61),        # драйверы, БДР, БДДС помесячно
    "дашборд": ("00_Дашборд", 15, 43, 1, 14),     # итоги по сезонам
    "проверки": ("13_Проверки", 1, 74, 1, 7),
    "сеть": ("06_Сеть_и_дети", 1, 40, 2, 61),
    "мерч": ("07_Мерч", 1, 32, 2, 61),
    "фк": ("10_Выплаты_ФК", 1, 37, 2, 61),
}


class RecalcError(RuntimeError):
    pass


def apply_overrides(src: Path, dst: Path, overrides: dict):
    """Копия книги с подставленными вводными. Кэш значений openpyxl не пишет,
    поэтому LibreOffice обязана пересчитать всё сама."""
    wb = openpyxl.load_workbook(src, data_only=False)
    applied = []
    for target, value in overrides.items():
        if "!" not in target:
            raise SystemExit(f"вводная должна быть вида Лист!Ячейка, а не {target!r}")
        sheet, addr = target.rsplit("!", 1)
        sheet = sheet.strip("'")
        if sheet not in wb.sheetnames:
            raise SystemExit(f"нет листа {sheet!r}")
        cell = wb[sheet][addr]
        if isinstance(cell.value, str) and cell.value.startswith("="):
            raise SystemExit(
                f"{target} — формула, а не вводная. Подставлять можно только вводные.")
        was = cell.value
        cell.value = value
        applied.append((target, was, value))
    wb.save(dst)
    return applied


def recalc(path: Path, workdir: Path) -> Path:
    """Пересчёт книги через LibreOffice. Возвращает путь к пересчитанной копии."""
    outdir = workdir / "out"
    outdir.mkdir(exist_ok=True)
    cmd = ["soffice", "--headless", "--norestore",
           f"-env:UserInstallation=file://{LO_PROFILE}",
           "--convert-to", "xlsx", "--outdir", str(outdir), str(path)]
    res = subprocess.run(cmd, capture_output=True, text=True, timeout=1800)
    out = outdir / path.name
    if not out.exists():
        raise RecalcError(
            "LibreOffice не пересчитала книгу.\n"
            f"stdout: {res.stdout.strip()}\nstderr: {res.stderr.strip()}")
    return out


def grab(path: Path):
    """Срез пересчитанной книги: подпись строки + значения по месяцам."""
    wb = openpyxl.load_workbook(path, data_only=True)
    result, errors = {}, []
    for block, (sname, r1, r2, c1, c2) in GRAB.items():
        if sname not in wb.sheetnames:
            continue
        ws = wb[sname]
        rows = {}
        for r in range(r1, min(r2, ws.max_row) + 1):
            label = ""
            for c in (1, 2):
                v = ws.cell(row=r, column=c).value
                if isinstance(v, str) and v.strip():
                    label = v.strip()
                    break
            vals = []
            for c in range(c1, min(c2, ws.max_column) + 1):
                v = ws.cell(row=r, column=c).value
                if isinstance(v, str) and v.startswith("#"):
                    errors.append(f"{sname}!{get_column_letter(c)}{r} = {v}")
                    v = None
                elif hasattr(v, "isoformat"):
                    v = v.isoformat()
                vals.append(v)
            if label or any(v is not None for v in vals):
                rows[str(r)] = {"label": label, "v": vals, "c0": c1}
        result[block] = {"sheet": sname, "rows": rows}
    return result, errors


def find_row(snapshot, block, label_part):
    """Найти строку среза по куску подписи."""
    rows = snapshot["blocks"].get(block, {}).get("rows", {})
    for r, rec in sorted(rows.items(), key=lambda kv: int(kv[0])):
        if label_part.lower() in (rec["label"] or "").lower():
            return r, rec
    return None, None


def nums(rec):
    return [v for v in rec["v"] if isinstance(v, (int, float))]


def summarize(snapshot):
    out = []
    for block, label in (("модель", "Выручка — всего"), ("модель", "EBITDA"),
                         ("модель", "Денежный остаток на конец"),
                         ("модель", "Всего активных детей"),
                         ("модель", "Новые школы — активные школы")):
        r, rec = find_row(snapshot, block, label)
        if not rec:
            continue
        v = nums(rec)
        if not v:
            continue
        if "остаток" in label or "дети" in label.lower() or "школ" in label:
            out.append(f"{rec['label']:<34} конец: {v[-1]:>18,.0f}   мин: {min(v):>15,.0f}")
        else:
            out.append(f"{rec['label']:<34} сумма: {sum(v):>18,.0f}   посл.: {v[-1]:>13,.0f}")
    return "\n".join(s.replace(",", " ") for s in out)


def checks_summary(snapshot):
    rows = snapshot["blocks"].get("проверки", {}).get("rows", {})
    bad = []
    for r, rec in sorted(rows.items(), key=lambda kv: int(kv[0])):
        for v in rec["v"]:
            if isinstance(v, str) and re.search(r"провал|ПРОВЕРИТЬ|FAIL|ошибк", v, re.I):
                bad.append(f"   строка {r}: {rec['label'][:70]} → {v}")
                break
    return bad


def read_inputs(path: Path):
    """Значения ключевых вводных из пересчитанной книги."""
    wb = openpyxl.load_workbook(path, data_only=True)
    out = {}
    for name, (sheet, addr) in KEY_INPUTS.items():
        if sheet in wb.sheetnames:
            v = wb[sheet][addr].value
            out[name] = v if isinstance(v, (int, float, str)) else str(v)
    return out


def read_scenario_no(path: Path):
    """Какой сценарий реально посчитался (по 01_Вводные!B194)."""
    wb = openpyxl.load_workbook(path, data_only=True)
    sheet, addr = SCENARIO_NO_CELL
    return wb[sheet][addr].value if sheet in wb.sheetnames else None


def do_run(model: Path, overrides: dict, tag: str, keep: bool = False):
    RUNS.mkdir(parents=True, exist_ok=True)
    BUILD.mkdir(exist_ok=True)
    started = datetime.now()
    with tempfile.TemporaryDirectory(prefix="spartak_run_") as td:
        work = Path(td)
        staged = work / model.name
        applied = apply_overrides(model, staged, overrides)
        for target, was, now in applied:
            print(f"вводная  {target}: {was!r} -> {now!r}")
        print("пересчёт LibreOffice...", flush=True)
        done = recalc(staged, work)
        blocks, errors = grab(done)
        scenario_no = read_scenario_no(done)
        inputs = read_inputs(done)
        kept = None
        if keep:
            RECALC.mkdir(parents=True, exist_ok=True)
            kept = RECALC / f"{tag}.xlsx"
            shutil.copy2(done, kept)

    snapshot = {
        "tag": tag,
        "scenario_no": scenario_no,
        "inputs": inputs,
        "model": str(model.relative_to(ROOT)),
        "overrides": {k: v for k, v in overrides.items()},
        "started": started.isoformat(timespec="seconds"),
        "seconds": round((datetime.now() - started).total_seconds(), 1),
        "cell_errors": errors,
        "blocks": blocks,
    }
    path = RUNS / f"{tag}.json"
    path.write_text(json.dumps(snapshot, ensure_ascii=False), encoding="utf-8")

    want = overrides.get("expect_scenario_no")
    names = {1: "негативный", 2: "базовый", 3: "позитивный"}
    print(f"\n--- прогон «{tag}» за {snapshot['seconds']} с ---")
    print(f"посчитан сценарий                  : "
          f"{names.get(scenario_no, '?')} (номер {scenario_no})")
    print(f"закрытие школ {inputs.get('закрытие_в_год', 0):.0%} в год, "
          f"запас {inputs.get('целевой_запас_мес', '?')} мес., "
          f"каникулы роялти {inputs.get('каникулы_роялти_мес', '?')} мес.")
    print(summarize(snapshot))
    if errors:
        print(f"\nОШИБКИ В ЯЧЕЙКАХ: {len(errors)}")
        for e in errors[:15]:
            print(f"   {e}")
    bad = checks_summary(snapshot)
    if bad:
        print(f"\nЛИСТ ПРОВЕРОК — не пройдено ({len(bad)}):")
        for b in bad[:20]:
            print(b)
    if kept:
        print(f"пересчитанная книга: {kept.relative_to(ROOT)}")
    print(f"срез: {path.relative_to(ROOT)}")
    return snapshot


def do_compare(a: str, b: str):
    pa, pb = RUNS / f"{a}.json", RUNS / f"{b}.json"
    for p in (pa, pb):
        if not p.exists():
            sys.exit(f"нет прогона {p.name}. Сначала запустите его с --tag")
    sa = json.loads(pa.read_text(encoding="utf-8"))
    sb = json.loads(pb.read_text(encoding="utf-8"))
    print(f"{'показатель':<38}{a:>18}{b:>18}{'дельта':>18}")
    for block, blk in sa["blocks"].items():
        rb = sb["blocks"].get(block, {}).get("rows", {})
        for r, rec in sorted(blk["rows"].items(), key=lambda kv: int(kv[0])):
            other = rb.get(r)
            if not other:
                continue
            va, vb = nums(rec), nums(other)
            if not va or not vb:
                continue
            sa_, sb_ = sum(va), sum(vb)
            if abs(sa_ - sb_) < max(1.0, abs(sa_) * 1e-9):
                continue
            print(f"{(rec['label'] or block+'!'+r)[:37]:<38}"
                  f"{sa_:>18,.0f}{sb_:>18,.0f}{sb_-sa_:>18,.0f}".replace(",", " "))


def main():
    ap = argparse.ArgumentParser(description="Прогон финмодели «Спартак»")
    ap.add_argument("--model", default=str(DEFAULT_MODEL))
    ap.add_argument("--set", action="append", default=[],
                    metavar="Лист!Ячейка=знач", help="подставить вводную")
    ap.add_argument("--scenario", choices=list(SCENARIOS), help="переключить сценарий")
    ap.add_argument("--tag", help="имя прогона (файл build/runs/<тег>.json)")
    ap.add_argument("--keep", action="store_true",
                    help="сохранить пересчитанную книгу в build/recalc/<тег>.xlsx "
                         "— в ней, в отличие от исходника, есть значения формул")
    ap.add_argument("--compare", nargs=2, metavar=("A", "B"), help="сравнить два прогона")
    a = ap.parse_args()

    if a.compare:
        return do_compare(*a.compare)

    model = Path(a.model)
    if not model.is_absolute():
        model = ROOT / model

    overrides = {}
    if a.scenario:
        overrides[SCENARIO_CELL] = SCENARIOS[a.scenario][0]
    for item in a.set:
        if "=" not in item:
            sys.exit(f"--set нужно вида Лист!Ячейка=значение, получено {item!r}")
        k, v = item.split("=", 1)
        try:
            v = int(v) if re.fullmatch(r"-?\d+", v.strip()) else float(v.replace(",", "."))
        except ValueError:
            pass
        overrides[k.strip()] = v

    tag = a.tag or (a.scenario if a.scenario else "база")
    snap = do_run(model, overrides, tag, keep=a.keep)
    if a.scenario:
        want = SCENARIOS[a.scenario][1]
        got = snap.get("scenario_no")
        if got != want:
            sys.exit(f"\nОШИБКА: просили сценарий «{a.scenario}» (номер {want}), "
                     f"а модель посчитала номер {got}. Прогон недостоверен.")


if __name__ == "__main__":
    main()
